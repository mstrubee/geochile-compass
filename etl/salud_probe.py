#!/usr/bin/env python3
"""Sondeo de SOLO LECTURA de las fuentes de «Inteligencia territorial de salud» (Fase 0).

No escribe en ninguna base. Hace HEAD/GET (con `Range` para los archivos grandes) y guarda un informe en una carpeta
local (por defecto `salud_probe_out/<fecha>/`, ignorada por git). Sirve para:
  * re-comprobar, en TU máquina, lo que la sesión de la Fase 0 no pudo alcanzar (Fonasa, Geofabrik, parte de DEIS);
  * medir columnas, filas, períodos, licencias y calidad con los archivos reales;
  * contar las filas de la base de manzanas del Censo 2024 (`--manzana-file`), que decide la Fase 2b.

Reglas: identidad honesta (User-Agent propio), SIN cabeceras de navegador, SIN reintentar bloqueos. Si un sitio responde con un
desafío de WAF («Just a moment…»), un 403 o corta la conexión, se registra tal cual y esa fuente queda en el plan B.
Los hosts `*_BASE` se pueden redirigir por variables de entorno solo para probar este script contra un servidor local.

Uso:  python etl/salud_probe.py [--out DIR] [--only est,rem20,egresos,isapre,casen,censo,fonasa,geofabrik]
                                [--years 2019-2025] [--full] [--manzana-file ARCHIVO] [--pause 0.3]
Contiene datos públicos y ninguna credencial.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import struct
import sys
import time
import unicodedata
import warnings
import zlib
from datetime import datetime, timezone
from pathlib import Path

import requests

sys.dont_write_bytecode = True

ROOT = Path(__file__).resolve().parents[1]
UA = "geochile-compass-salud-probe/1 (solo lectura; contacto: equipo GeoChile)"
SAMPLE_BYTES = 1_500_000          # lo que se pide por Range de cada zip (~1,5 MB comprimidos)
TAIL_BYTES = 70_000               # cola del zip donde vive el directorio central


def base(name: str) -> str:
    """URL base de cada host; se puede redirigir con PROBE_<NOMBRE> (tests)."""
    d = {
        "CKAN": "https://datos.gob.cl/api/3/action",
        "DEIS_REPO": "https://repositoriodeis.minsal.cl",
        "SUPERSALUD": "https://www.superdesalud.gob.cl",
        "CASEN": "https://observatorio.ministeriodesarrollosocial.gob.cl",
        "INE": "https://www.ine.gob.cl",
        "CENSO": "https://censo2024.ine.gob.cl",
        "REDATAM": "https://redatam.ine.gob.cl",
        "FONASA": "https://www.fonasa.cl",
        "GEOFABRIK": "https://download.geofabrik.de",
    }
    return os.environ.get(f"PROBE_{name}", d[name]).rstrip("/")


# ---------------------------------------------------------------------------------------------------------------------------
# Cliente: clasifica cada resultado; NO reintenta bloqueos
# ---------------------------------------------------------------------------------------------------------------------------
def classify(resp: requests.Response | None, exc: Exception | None) -> dict:
    """kind ∈ ok | waf_challenge | site_403 | site_4xx | site_5xx | proxy_policy_denied | reset | dns | tls | timeout | error."""
    if exc is not None:
        m = str(exc)
        if "CONNECT tunnel failed" in m or "Tunnel connection failed" in m or "403 Forbidden" in m and "proxy" in m.lower():
            return {"kind": "proxy_policy_denied", "detail": m[:200]}
        if "Connection reset" in m or "Connection aborted" in m or "RemoteDisconnected" in m:
            return {"kind": "reset", "detail": m[:200]}
        if "Name or service not known" in m or "Could not resolve" in m or "NameResolutionError" in m:
            return {"kind": "dns", "detail": m[:200]}
        if "CERTIFICATE" in m.upper() or "SSL" in m.upper():
            return {"kind": "tls", "detail": m[:200]}
        if isinstance(exc, requests.Timeout):
            return {"kind": "timeout", "detail": m[:200]}
        return {"kind": "error", "detail": m[:200]}
    code = resp.status_code
    if code in (403, 429, 503) and (b"Just a moment" in resp.content[:4000] or resp.headers.get("cf-mitigated") == "challenge"):
        return {"kind": "waf_challenge", "status": code, "detail": "desafío de WAF (Cloudflare); no se elude"}
    if code == 403:
        return {"kind": "site_403", "status": code, "server": resp.headers.get("server", "")}
    if 200 <= code < 400 or code == 206:
        return {"kind": "ok", "status": code}
    return {"kind": "site_5xx" if code >= 500 else "site_4xx", "status": code}


class Client:
    def __init__(self, pause: float = 0.3, timeout: int = 60):
        self.s = requests.Session()
        self.s.headers["User-Agent"] = UA
        self.pause, self.timeout = pause, timeout
        self.log: list[dict] = []

    def req(self, method: str, url: str, **kw) -> tuple[requests.Response | None, dict]:
        time.sleep(self.pause)
        t0 = time.time()
        try:
            r = self.s.request(method, url, timeout=self.timeout, allow_redirects=True, **kw)
            c, exc = classify(r, None), None
        except requests.RequestException as e:  # noqa: PERF203
            r, c, exc = None, classify(None, e), e
        entry = {"method": method, "url": url, "seconds": round(time.time() - t0, 2), **c}
        if r is not None:
            entry["content_type"] = r.headers.get("content-type", "")
            entry["content_length"] = r.headers.get("content-length", "")
            entry["last_modified"] = r.headers.get("last-modified", "")
        self.log.append(entry)
        return r, entry

    def head(self, url):
        return self.req("HEAD", url)

    def get(self, url, **kw):
        return self.req("GET", url, **kw)

    def range(self, url, start: int, end: int):
        return self.req("GET", url, headers={"Range": f"bytes={start}-{end}"})


# ---------------------------------------------------------------------------------------------------------------------------
# Utilidades puras (se prueban sin red)
# ---------------------------------------------------------------------------------------------------------------------------
def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def decode_bytes(b: bytes) -> tuple[str, str]:
    """UTF-8 si se puede; si no, cp1252 (los archivos del DEIS y de la Superintendencia no son UTF-8)."""
    try:
        return b.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return b.decode("cp1252", errors="replace"), "cp1252"


def read_zip_directory(tail: bytes, file_size: int) -> list[dict]:
    """Directorio central de un zip a partir de su cola. Devuelve nombre, método, tamaños y offset del encabezado local."""
    i = tail.rfind(b"PK\x05\x06")
    if i < 0:
        raise ValueError("no hay EOCD en la cola")
    n, cd_size, cd_off = struct.unpack("<HII", tail[i + 10:i + 20])
    if cd_off == 0xFFFFFFFF or cd_size == 0xFFFFFFFF:
        raise ValueError("zip64: no soportado por el sondeo")
    start = len(tail) - (file_size - cd_off)
    if start < 0:
        raise ValueError("el directorio central no cabe en la cola pedida")
    cd, out, p = tail[start:], [], 0
    for _ in range(n):
        if cd[p:p + 4] != b"PK\x01\x02":
            raise ValueError("directorio central corrupto")
        meth, = struct.unpack("<H", cd[p + 10:p + 12])
        csz, usz = struct.unpack("<II", cd[p + 20:p + 28])
        nlen, elen, clen = struct.unpack("<HHH", cd[p + 28:p + 34])
        off, = struct.unpack("<I", cd[p + 42:p + 46])
        name = cd[p + 46:p + 46 + nlen].decode("cp437", errors="replace")
        out.append({"name": name, "method": meth, "compressed": csz, "uncompressed": usz, "offset": off})
        p += 46 + nlen + elen + clen
    return out


def inflate_entry_head(data: bytes) -> tuple[str, bytes]:
    """`data` empieza en un encabezado local de zip (PK\\x03\\x04). Descomprime lo que alcance (deflate parcial)."""
    if data[:4] != b"PK\x03\x04":
        raise ValueError("no empieza en un encabezado local de zip")
    meth, = struct.unpack("<H", data[8:10])
    nlen, elen = struct.unpack("<HH", data[26:30])
    name = data[30:30 + nlen].decode("cp437", errors="replace")
    body = data[30 + nlen + elen:]
    if meth == 0:
        return name, body
    if meth != 8:
        raise ValueError(f"método de compresión {meth} no soportado")
    return name, zlib.decompressobj(-15).decompress(body)


def profile_rows(text: str, delimiter: str, *, header: bool = True, comuna_col: str | None = None, ref_cuts: set[str] | None = None) -> dict:
    """Perfil de las filas completas de una muestra (descarta la última, que puede venir cortada)."""
    lines = text.splitlines()[:-1] if text and not text.endswith("\n") else text.splitlines()
    cols = lines[0].split(delimiter) if header and lines else []
    rows = [l.split(delimiter) for l in (lines[1:] if header else lines)]
    rep = {"columns": cols, "n_columns": len(cols) or (len(rows[0]) if rows else 0), "sample_rows": len(rows), "header": header}
    if comuna_col and comuna_col in cols:
        i = cols.index(comuna_col)
        vals = [r[i] for r in rows if len(r) > i]
        star = sum(1 for v in vals if v == "*")
        rep["comuna_masked_asterisk"] = {"rows": star, "ratio": round(star / len(vals), 4) if vals else None}
        good = [v for v in vals if v != "*"]
        rep["comuna_code_lengths"] = {str(k): sum(1 for v in good if len(v) == k) for k in sorted({len(v) for v in good})}
        if ref_cuts is not None:
            bad = {}
            for v in good:
                if v not in ref_cuts:
                    bad[v] = bad.get(v, 0) + 1
            rep["comuna_codes_not_in_reference"] = dict(sorted(bad.items(), key=lambda kv: -kv[1])[:10])
        rep["first_comunas"] = good[:5]
        rep["looks_sorted_by_comuna"] = good[:200] == sorted(good[:200]) and len(set(good[:200])) < 20 if good else None
    return rep


# Isapre (5 años) -> grupos de egreso DEIS (10 años). «menor de un año» + «1 a 9» = 0-9 y solo se comparan juntos.
ISAPRE_TO_EGRESO = {
    "0-4": "0 a 9", "5-9": "0 a 9", "10-14": "10 a 19", "15-19": "10 a 19", "20-24": "20 a 29", "25-29": "20 a 29",
    "30-34": "30 a 39", "35-39": "30 a 39", "40-44": "40 a 49", "45-49": "40 a 49", "50-54": "50 a 59", "55-59": "50 a 59",
    "60-64": "60 a 69", "65-69": "60 a 69", "70-74": "70 a 79", "75-79": "70 a 79", "80-84": "80 a 89", "85-89": "80 a 89",
    "90-94": "90 y más", "95-99": "90 y más", ">=100": "90 y más",
}


def ref_cuts() -> tuple[set[str], dict[str, list[str]]]:
    p = ROOT / "public" / "codigos_territoriales.csv"
    if not p.exists():
        return set(), {}
    by: dict[str, list[str]] = {}
    for r in csv.DictReader(open(p, encoding="utf-8")):
        by.setdefault(norm(r["commune_name"]), []).append(r["commune_id"])
    return {c for v in by.values() for c in v}, by


# ---------------------------------------------------------------------------------------------------------------------------
# Fuentes
# ---------------------------------------------------------------------------------------------------------------------------
def ckan_package(c: Client, pid: str) -> dict:
    r, e = c.get(f"{base('CKAN')}/package_show?id={pid}")
    if r is None or e["kind"] != "ok":
        return {"access": e}
    d = r.json()["result"]
    return {"access": e, "title": d.get("title"), "license_id": d.get("license_id"), "metadata_modified": d.get("metadata_modified"),
            "organization": (d.get("organization") or {}).get("title"),
            "resources": [{"name": x.get("name"), "format": x.get("format"), "size": x.get("size"), "url": x.get("url"),
                           "last_modified": x.get("last_modified") or x.get("metadata_modified")} for x in d.get("resources", [])]}


def probe_est(c: Client, out: Path, full: bool) -> dict:
    pk = ckan_package(c, "establecimientos-de-salud-vigentes")
    res = next((x for x in pk.get("resources", []) if (x.get("format") or "").lower() == "csv"), None)
    rep = {"package": pk}
    if not res:
        return rep
    r, e = c.get(res["url"]) if full else c.range(res["url"], 0, 250_000)
    rep["download"] = e
    if r is None or e["kind"] != "ok":
        return rep
    text, enc = decode_bytes(r.content)
    cuts, _ = ref_cuts()
    prof = profile_rows(text, ";", comuna_col="ComunaCodigo", ref_cuts=cuts)
    rep.update({"encoding": enc, "sha256": hashlib.sha256(r.content).hexdigest() if full else None, "complete_file": full, **prof})
    if full:
        rows = list(csv.DictReader(io.StringIO(text), delimiter=";"))
        rep["rows"] = len(rows)
        for k in ("EstadoFuncionamiento", "NivelComplejidadEstabGlosa", "TipoAtencionEstabGlosa", "DependenciaAdministrativa"):
            cnt: dict[str, int] = {}
            for x in rows:
                cnt[x[k]] = cnt.get(x[k], 0) + 1
            rep[k] = dict(sorted(cnt.items(), key=lambda kv: -kv[1]))
        miss = sum(1 for x in rows if not x["Latitud"].strip() or not x["Longitud"].strip())
        rep["sin_coordenadas"] = {"rows": miss, "ratio": round(miss / len(rows), 4) if rows else None}
        rep["comunas_distintas"] = len({x["ComunaCodigo"] for x in rows})
    return rep


def probe_rem20(c: Client, out: Path, full: bool) -> dict:
    pk = ckan_package(c, "indicadores-hospitalarios-rem20")
    res = next((x for x in pk.get("resources", []) if (x.get("format") or "").lower() == "csv"), None)
    rep = {"package": pk}
    if res:
        r, e = c.range(res["url"], 0, 300_000)
        rep["download"] = e
        if r is not None and e["kind"] == "ok":
            text, enc = decode_bytes(r.content)
            rep.update({"encoding": enc, **profile_rows(text, ";")})
    return rep


def probe_egresos(c: Client, out: Path, years: list[int]) -> dict:
    """Egresos DEIS: patrón de URL OBSERVADO (no es un enlace oficial verificado): <repo>/DatosAbiertos/EGRESOS/EGRESOS_<año>.zip."""
    cuts, _ = ref_cuts()
    rep = {"url_pattern": f"{base('DEIS_REPO')}/DatosAbiertos/EGRESOS/EGRESOS_<año>.zip", "url_pattern_status": "[NO COMPROBADO como enlace oficial]",
           "years": {}}
    for y in years:
        url = f"{base('DEIS_REPO')}/DatosAbiertos/EGRESOS/EGRESOS_{y}.zip"
        item: dict = {"url": url}
        r, e = c.head(url)
        item["head"] = e
        rep["years"][y] = item
        if r is None or e["kind"] != "ok" or "zip" not in e.get("content_type", ""):
            continue
        size = int(e["content_length"] or 0)
        if not size:
            continue
        try:
            tr, te = c.range(url, max(0, size - TAIL_BYTES), size - 1)
            entries = read_zip_directory(tr.content, size) if tr is not None and te["kind"] == "ok" else []
            item["entries"] = [{k: v for k, v in x.items() if k != "offset"} for x in entries]
            data_entries = [x for x in entries if x["name"].lower().endswith((".csv", ".txt"))]
            if not data_entries:
                continue
            ent = max(data_entries, key=lambda x: x["uncompressed"])
            hr, he = c.range(url, ent["offset"], ent["offset"] + SAMPLE_BYTES)
            if hr is None or he["kind"] != "ok":
                item["sample"] = he
                continue
            name, raw = inflate_entry_head(hr.content)
            text, enc = decode_bytes(raw)
            prof = profile_rows(text, ";", comuna_col="COMUNA_RESIDENCIA", ref_cuts=cuts)
            first = text.splitlines()[0] if text else ""
            item["sample"] = {"entry": name, "uncompressed_bytes": ent["uncompressed"], "encoding": enc,
                              "has_establishment_column": bool(re.search(r"ESTAB|HOSP", first.upper().replace("PERTENENCIA_ESTABLECIMIENTO", ""))),
                              "estimated_rows": round(ent["uncompressed"] / max(1, len(text) / max(1, prof["sample_rows"] + 1))), **prof}
        except (ValueError, zlib.error, struct.error) as ex:
            item["error"] = str(ex)[:200]
    return rep


def analyze_isapre_xlsx(path: Path) -> dict:
    import openpyxl
    cuts, by = ref_cuts()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rep: dict = {"sheets": wb.sheetnames, "sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest()}
        for sheet in ("Total Cotizantes", "Total Cargas"):
            if sheet not in wb.sheetnames:
                continue
            rows = list(wb[sheet].iter_rows(values_only=True))
            hdr = next((r for r in rows if r and "Total" in r and "0-4" in r), None)
            if hdr is None:
                continue
            tot_i = list(hdr).index("Total")
            com = [r for r in rows if r and isinstance(r[0], int) and r[3] not in (None, "-") and not str(r[1]).startswith("Total")]
            pais = next((r for r in rows if r and r[1] and str(r[1]).strip() == "Total Pais"), None)
            real = [r for r in com if not str(r[3]).strip().lower().startswith("sin dato")]
            unmatched = sorted({str(r[3]).strip() for r in real if norm(r[3]) not in by})
            rep[sheet] = {"age_bands": [x for x in hdr[4:] if x not in (None, "Total")], "rows_comuna_o_sin_dato": len(com),
                          "sum_comunas": sum(r[tot_i] for r in com), "total_pais_publicado": pais[tot_i] if pais else None,
                          "cuadra": bool(pais and sum(r[tot_i] for r in com) == pais[tot_i]),
                          "sin_dato_comuna": sum(r[tot_i] for r in com if str(r[3]).strip().lower().startswith("sin dato")),
                          "nombres_sin_calzar_exacto": unmatched,
                          "comunas_de_referencia_ausentes": sorted(set(by) - {norm(r[3]) for r in real})}
    return rep


def probe_isapre(c: Client, out: Path, full: bool) -> dict:
    sup = base("SUPERSALUD")
    rep: dict = {"comunal": {}}
    for y, path in ((2025, "2026/01/estadistica-anual-de-cartera-de-beneficiarios-comunal-2025.xlsx"),
                    (2024, "2025/01/estadistica-anual-de-cartera-de-beneficiarios-comunal-2024.xlsx")):
        url = f"{sup}/app/uploads/{path}"
        r, e = c.get(url)
        item = {"url": url, "access": e}
        if r is not None and e["kind"] == "ok":
            f = out / f"isapre_comunal_{y}.xlsx"
            f.write_bytes(r.content)
            try:
                item.update(analyze_isapre_xlsx(f))
            finally:
                f.unlink(missing_ok=True)       # el archivo descargado no se conserva; solo su hash y el análisis
        rep["comunal"][y] = item
    zurl = "https://datos.superdesalud.gob.cl/archivos/Beneficiarios/2026/Beneficiarios_2026t2.zip"
    rep["trimestral_txt"] = {"url": zurl, "nota": "HEAD no es fiable en este host: se usa GET con Range"}
    r, e = c.range(zurl, 0, SAMPLE_BYTES)
    rep["trimestral_txt"]["access"] = e
    if r is not None and e["kind"] == "ok" and r.content[:4] == b"PK\x03\x04":
        name, raw = inflate_entry_head(r.content)
        text, enc = decode_bytes(raw)
        first = text.splitlines()[0] if text else ""
        rep["trimestral_txt"].update({"entry": name, "encoding": enc, "delimiter": "|", "first_line_is_header": not bool(re.match(r"^\d{4}t\d", first)),
                                      "n_columns": len(first.split("|")), "diccionario": "[NO COMPROBADO] vive en un modal por JS de la página"})
    return rep


def probe_casen(c: Client, out: Path, full: bool) -> dict:
    root = f"{base('CASEN')}/storage/docs/casen/2024"
    rep: dict = {"files": {}}
    for f in ("casen_2024_provincia_comuna.sav", "casen_2024.sav", "Libro_de_codigos_Casen_2024_prov_comuna.xlsx", "Nota_uso_bases_de_datos_Casen_2024.pdf"):
        r, e = c.head(f"{root}/{f}")
        rep["files"][f] = e
    r, e = c.get(f"{root}/Libro_de_codigos_Casen_2024_prov_comuna.xlsx")
    if r is not None and e["kind"] == "ok":
        import openpyxl
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            rows = [[x for x in row if x not in (None, "")] for row in ws.iter_rows(values_only=True)]
        rep["variables_base_comunal"] = [row[0] for row in rows if len(row) >= 3 and row[0] not in ("Nombre variable",)][:30]
        rep["nota"] = "la base provincia-comuna solo trae claves y factores; la previsión vive en casen_2024.sav (nacional)"
    return rep


def probe_censo(c: Client, out: Path, manzana_file: str | None) -> dict:
    rep: dict = {"hosts": {}}
    for k, url in (("censo2024", f"{base('CENSO')}/resultados/"), ("redatam", f"{base('REDATAM')}/redbin/RpWebEngine.exe/Portal?BASE=CENSO_2024&lang=esp"),
                   ("ine_licencia", f"{base('INE')}/terminos-de-uso-y-licencia-de-datos-abiertos")):
        rep["hosts"][k] = c.head(url)[1]
    rep["comunicados"] = ["INE 2025-12-04: base comuna-área (microdatos)", "INE 2025-12-16: base manzana-entidad (189 variables) + cartografía"]
    rep["urls_de_descarga"] = "[NO COMPROBADO] los enlaces son dinámicos; descargarlos desde censo2024.ine.gob.cl/resultados y pasar el archivo con --manzana-file"
    if manzana_file:
        p = Path(manzana_file)
        enc_text = p.read_bytes()[:200_000]
        text, enc = decode_bytes(enc_text)
        first = text.splitlines()[0] if text else ""
        delim = max(";,|\t", key=first.count)
        n = sum(1 for _ in open(p, "rb")) - 1
        rep["manzana_file"] = {"path_name": p.name, "bytes": p.stat().st_size, "encoding": enc, "delimiter": delim, "columns": first.split(delim), "rows": n,
                               "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
        rep["decision_fase_2b"] = "filas = manzanas/entidades; si superan ~50.000 no se pinta GeoJSON en Leaflet: Fase 2b (MVT/PMTiles)"
    return rep


def probe_fonasa(c: Client, out: Path) -> dict:
    """Solo alcanzabilidad. Un 403 del sitio (WAF) NO se elude: se anota y se pide la descarga manual."""
    return {"hosts": {u: c.head(u)[1] for u in (f"{base('FONASA')}/", "https://datosabiertos.fonasa.cl/")},
            "premisa": "Fonasa: comuna = comuna del establecimiento APS de inscripción [NO COMPROBADO]",
            "plan_b": "descargar a mano desde la sección de estadísticas/datos abiertos de Fonasa y pasar el archivo a la Fase 1b"}


def probe_geofabrik(c: Client, out: Path) -> dict:
    return {"chile_pbf": c.head(f"{base('GEOFABRIK')}/south-america/chile-latest.osm.pbf")[1]}


# ---------------------------------------------------------------------------------------------------------------------------
def write_report(out: Path, report: dict) -> None:
    out.mkdir(parents=True, exist_ok=True)
    (out / "report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    lines = ["# Resumen del sondeo de salud", "", f"Generado: {report['generated_at']}", "", "| Fuente | Resultado |", "|---|---|"]
    kinds: dict[str, set] = {}
    for e in report["http_log"]:
        host = re.sub(r"^https?://([^/]+).*", r"\1", e["url"])
        kinds.setdefault(host, set()).add(e["kind"] + (f" {e.get('status')}" if e.get("status") else ""))
    for h, k in sorted(kinds.items()):
        lines.append(f"| {h} | {', '.join(sorted(k))} |")
    (out / "RESUMEN.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def parse_years(s: str) -> list[int]:
    a, _, b = s.partition("-")
    return list(range(int(a), int(b or a) + 1))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--out", default=os.environ.get("SALUD_PROBE_OUT") or str(ROOT / "salud_probe_out" / datetime.now(timezone.utc).strftime("%Y-%m-%d")))
    ap.add_argument("--only", default="est,rem20,egresos,isapre,casen,censo,fonasa,geofabrik")
    ap.add_argument("--years", default="2019-2025")
    ap.add_argument("--full", action="store_true", help="descarga completo el CSV de establecimientos (2,5 MB) para medir y hashear")
    ap.add_argument("--manzana-file", help="CSV/TXT de la base manzana-entidad del Censo 2024 para contar filas (decide la Fase 2b)")
    ap.add_argument("--pause", type=float, default=0.3)
    a = ap.parse_args(argv)
    out, only = Path(a.out), set(a.only.split(","))
    out.mkdir(parents=True, exist_ok=True)
    c = Client(a.pause)
    report: dict = {"generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "ua": UA}
    steps = {"est": lambda: probe_est(c, out, a.full), "rem20": lambda: probe_rem20(c, out, a.full),
             "egresos": lambda: probe_egresos(c, out, parse_years(a.years)), "isapre": lambda: probe_isapre(c, out, a.full),
             "casen": lambda: probe_casen(c, out, a.full), "censo": lambda: probe_censo(c, out, a.manzana_file),
             "fonasa": lambda: probe_fonasa(c, out), "geofabrik": lambda: probe_geofabrik(c, out)}
    for k, fn in steps.items():
        if k in only:
            try:
                report[k] = fn()
            except Exception as ex:  # noqa: BLE001 - una fuente caída no debe tumbar el resto del sondeo
                report[k] = {"error": f"{type(ex).__name__}: {ex}"[:300]}
            print(f"[{k}] listo", file=sys.stderr)
    report["http_log"] = c.log
    write_report(out, report)
    print(f"Informe en {out}/report.json y RESUMEN.md")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
