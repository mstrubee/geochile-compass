"""Lector de los XLSX de MINVU CEDOC (permisos de edificación: viviendas por comuna).

Lee SOLO valores (nunca modifica el archivo) y se guía por la ESTRUCTURA, no por posiciones fijas:

  fila de encabezado : la que trae "Región" | "Comuna" | años      (en los archivos reales, la 6)
  fila de meses      : la siguiente ("Ene".."Dic" bajo el año)     (la 7)
  "Total País"       : primera fila de datos                        (la 8)
  subtotales regional: Región con Comuna vacía
  filas de comuna    : Comuna con texto
  notas al pie       : desde la celda de Comuna que empieza con "Notas" -> se DETIENE el parseo de datos

Estructura COMPROBADA el 2026-10-16 sobre los dos archivos publicados el 2026-09-11
(SHA-256 en etl/permisos/tests/fixtures/README.md). `max_row` no es confiable: la hoja m2_casas del
archivo mensual declara 65.536 filas por formato basura, por eso se lee de forma perezosa y se corta en las notas.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from typing import Any

import openpyxl

from .norm import MESES, norm

EXPECTED_SHEETS = (
    "número_total", "m2_total",
    "número_departamentos", "m2_departamentos",
    "número_casas", "m2_casas",
)
NOTES_TAIL_ROWS = 12          # filas que se leen después de "Notas" para juntar el texto de las notas
HEADER_SEARCH_ROWS = 30


class LayoutError(Exception):
    """El archivo no tiene la estructura esperada. `code` calza con permit_rejected_rows.reason_code."""

    def __init__(self, message: str, code: str = "header_unrecognized"):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ColSpec:
    col: int                   # columna 1-based (A=1)
    year: int
    month: int | None          # None = columna anual
    header_text: str           # texto crudo del encabezado del año (p. ej. "2024  (*)")
    mark: str | None           # "(*)", "(**)", "(***)" o None

    @property
    def granularity(self) -> str:
        return "year" if self.month is None else "month"

    @property
    def period(self) -> date:
        return date(self.year, self.month or 1, 1)

    @property
    def key(self) -> str:
        return f"{self.period.isoformat()}:{self.granularity}"


@dataclass
class RawRow:
    row: int                   # fila 1-based en la hoja
    region_label: str
    comuna_label: str
    values: dict[int, Any]     # col -> valor crudo


@dataclass
class Note:
    row: int
    marker: str | None
    text: str


@dataclass
class SheetData:
    name: str
    header_row: int
    cols: list[ColSpec]
    total_row: int
    total_values: dict[int, Any]
    region_rows: list[RawRow]
    rows: list[RawRow]
    notes: list[Note] = field(default_factory=list)


def _s(v: Any) -> str:
    return "" if v is None else str(v)


def _parse_header(header: tuple, months: tuple) -> list[ColSpec]:
    cols: list[ColSpec] = []
    year: int | None = None
    mark: str | None = None
    text = ""
    width = max(len(header), len(months))
    for idx in range(2, width):                         # desde la columna C
        h = header[idx] if idx < len(header) else None
        m = months[idx] if idx < len(months) else None
        if h is not None and _s(h).strip() != "":
            mt = re.match(r"^\s*(\d{4})\b", _s(h))
            if not mt:
                raise LayoutError(f"encabezado de año ilegible en la columna {idx + 1}: {h!r}")
            year = int(mt.group(1))
            mk = re.search(r"\((\*+)\)", _s(h))
            mark = f"({mk.group(1)})" if mk else None
            text = _s(h)
        if m is None or _s(m).strip() == "":
            if h is not None and _s(h).strip() != "":
                cols.append(ColSpec(idx + 1, year, None, text, mark))
            continue                                     # relleno: ni año ni mes
        month = MESES.get(norm(m))
        if month is None or year is None:
            raise LayoutError(f"mes ilegible en la columna {idx + 1}: {m!r}")
        cols.append(ColSpec(idx + 1, year, month, text, mark))
    if not cols:
        raise LayoutError("no se encontraron columnas de datos en el encabezado")
    seen = {(c.year, c.month) for c in cols}
    if len(seen) != len(cols):
        raise LayoutError("períodos duplicados en el encabezado")
    return cols


def read_sheet(ws, name: str) -> SheetData:
    header_row = None
    header: tuple = ()
    months: tuple = ()
    total_row = None
    total_values: dict[int, Any] = {}
    cols: list[ColSpec] = []
    region_rows: list[RawRow] = []
    rows: list[RawRow] = []
    notes: list[Note] = []
    in_notes_until = None

    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if in_notes_until is not None:
            if r > in_notes_until:
                break
            texts = [_s(v).strip() for v in row[2:] if isinstance(v, str) and v.strip()]
            marker = _s(row[1]).strip() if len(row) > 1 and row[1] is not None else None
            if texts or marker:
                notes.append(Note(r, marker or None, texts[0] if texts else ""))
            continue

        if header_row is None:
            if r > HEADER_SEARCH_ROWS:
                raise LayoutError(f"hoja {name!r}: no se encontró la fila 'Región | Comuna' en las primeras {HEADER_SEARCH_ROWS} filas")
            if len(row) >= 2 and norm(_s(row[0])) == "region" and norm(_s(row[1])) == "comuna":
                header_row = r
                header = row
            continue

        if r == header_row + 1:
            months = row
            cols = _parse_header(header, months)
            continue

        width = len(row)
        a = _s(row[0]).strip() if width > 0 else ""
        b = _s(row[1]).strip() if width > 1 else ""
        values = {c.col: (row[c.col - 1] if c.col - 1 < width else None) for c in cols}

        if b and norm(b).startswith("notas"):
            in_notes_until = r + NOTES_TAIL_ROWS
            texts = [_s(v).strip() for v in row[2:] if isinstance(v, str) and v.strip()]
            notes.append(Note(r, b, texts[0] if texts else ""))
            continue
        if a and norm(a) == "total pais":
            total_row, total_values = r, values
            continue
        if b:
            rows.append(RawRow(r, a, b, values))
        elif a:
            region_rows.append(RawRow(r, a, "", values))
        # filas vacías: se ignoran

    if header_row is None:
        raise LayoutError(f"hoja {name!r}: sin encabezado")
    if total_row is None:
        raise LayoutError(f"hoja {name!r}: sin fila 'Total País'", "header_unrecognized")
    if not rows:
        raise LayoutError(f"hoja {name!r}: sin filas de comuna")
    if not notes:
        raise LayoutError(f"hoja {name!r}: sin notas al pie (no se puede determinar qué años son definitivos)", "header_unrecognized")
    return SheetData(name, header_row, cols, total_row, total_values, region_rows, rows, notes)


def read_workbook(path: str) -> tuple[dict[str, SheetData], dict[str, Any]]:
    """Devuelve las 6 hojas parseadas y metadatos del libro (propiedades declaradas por el archivo)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = tuple(wb.sheetnames)
        if set(names) != set(EXPECTED_SHEETS):
            raise LayoutError(f"hojas inesperadas: {names} (se esperaban {EXPECTED_SHEETS})")
        sheets = {n: read_sheet(wb[n], n) for n in EXPECTED_SHEETS}
        props = wb.properties
        meta = {
            "creator": props.creator,
            "created": props.created.isoformat() if props.created else None,
            "modified": props.modified.isoformat() if props.modified else None,
        }
    finally:
        wb.close()
    return sheets, meta
