"""Pruebas de extremo a extremo contra una base REAL con las migraciones 1 y 2 aplicadas.

Solo corren si TEST_DATABASE_URL está definida Y el nombre de la base contiene 'test' (limpian las tablas del módulo):

  TEST_DATABASE_URL=postgresql://postgres:***@127.0.0.1/permtest python -m pytest etl/permisos/tests/test_db.py
"""
import os
import random
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from conftest import ANUAL, MENSUAL

URL = os.environ.get("TEST_DATABASE_URL")
pytestmark = pytest.mark.skipif(not URL, reason="define TEST_DATABASE_URL (base de PRUEBA con las migraciones aplicadas)")
ETL = Path(__file__).resolve().parents[2]


def cli(script, args, extra_env=None):
    env = dict(os.environ, SUPABASE_DB_URL=URL, **(extra_env or {}))
    return subprocess.run([sys.executable, "-B", str(ETL / script), *args], capture_output=True, text=True, env=env)


@pytest.fixture(scope="module")
def conn():
    import psycopg
    c = psycopg.connect(URL, autocommit=True)
    name = c.execute("select current_database()").fetchone()[0]
    if "test" not in name:
        pytest.skip(f"la base {name!r} no parece de prueba: no se toca")
    c.execute("""truncate public.building_permit_statistics, public.permit_stage, public.snapshot_control_totals,
                 public.permit_rejected_rows, public.source_snapshots, public.source_checks, public.source_alerts
                 restart identity cascade""")
    yield c
    c.close()


def q(conn, sql, *a):
    return conn.execute(sql, a).fetchall()


def test_seed_idempotente_y_check(conn):
    r1 = cli("permisos_comunas.py", ["seed"])
    assert r1.returncode == 0, r1.stderr
    r2 = cli("permisos_comunas.py", ["seed"])
    assert '"geometries_set": 0' in r2.stdout and '"geometries_unchanged": 346' in r2.stdout
    assert cli("permisos_comunas.py", ["check"]).returncode == 0
    st = dict(q(conn, "select geom_status, count(*) from public.comunas group by 1"))
    assert st == {"ok": 343, "repaired": 2, "absent": 1}
    assert q(conn, "select count(*) from public.comuna_geom_repairs")[0][0] == 2      # la 2.ª siembra no repitió reparaciones


def test_check_detecta_divergencia(conn):
    conn.execute("update public.comunas set name = 'Otro' where cut = '13101'")
    r = cli("permisos_comunas.py", ["check"])
    assert r.returncode == 1 and "13101" in r.stdout
    conn.execute("update public.comunas set name = 'Santiago' where cut = '13101'")
    assert cli("permisos_comunas.py", ["check"]).returncode == 0


def test_carga_anual_y_reintento(conn):
    r = cli("permisos_minvu.py", ["--file", str(ANUAL), "--origin-url", "https://example.test/anual.xlsx", "--no-storage", "--declared-date", "2026-09-11"])
    assert r.returncode == 0, r.stdout + r.stderr
    assert '"result": "published"' in r.stdout
    assert q(conn, "select count(*) from public.building_permit_statistics where to_snapshot is null")[0][0] == 346 * 6 * 6
    sums = dict(q(conn, """select extract(year from period)::int, sum(value) from public.building_permit_statistics b
                           join public.permit_series ps on ps.id = b.series_id where ps.source_sheet = 'número_total' group by 1"""))
    assert sums[2017] == 133914 and sums[2018] == 150491 and sums[2025] == 66947
    again = cli("permisos_minvu.py", ["--file", str(ANUAL), "--origin-url", "https://example.test/anual.xlsx", "--no-storage"])
    assert again.returncode == 0 and '"already_loaded"' in again.stdout
    assert q(conn, "select count(*) from public.source_snapshots")[0][0] == 1
    st = q(conn, "select status, declared_date::text, declared_date_note, rejected_count from public.source_snapshots")[0]
    assert st == ("loaded", "2026-09-11", "indicada por el operador (--declared-date)", 0)


def test_cada_valor_de_la_base_es_el_del_archivo(conn):
    """Compara celdas al azar entre la base y el XLSX leído de forma independiente (openpyxl directo)."""
    wb = openpyxl.load_workbook(ANUAL, data_only=True)
    ws = wb["m2_casas"]
    layout = q(conn, "select layout from public.source_snapshots order by n limit 1")[0][0]
    rnd = random.Random(7)
    checked = 0
    for cut in rnd.sample(sorted(layout["rows"]), 40):
        rows = layout["rows"][cut]
        for key in ("2002-01-01:year", "2018-01-01:year", "2025-01-01:year"):
            period = key.split(":")[0]
            row = rows[0] if isinstance(rows, list) else (rows["before"] if period < rows["split"] else rows["from"])
            expected = ws.cell(row, layout["cols"][key]).value
            got = q(conn, """select value from public.building_permit_statistics b join public.permit_series ps on ps.id = b.series_id
                             where ps.source_sheet = 'm2_casas' and b.comuna_cut = %s and b.period = %s and b.to_snapshot is null""", cut, period)
            assert got == [(expected,)], (cut, key, got, expected)
            checked += 1
    assert checked == 120


def test_carga_mensual_ventana_y_rpc(conn):
    r = cli("permisos_minvu.py", ["--file", str(MENSUAL), "--origin-url", "https://example.test/mensual.xlsx", "--no-storage"])
    assert r.returncode == 0, r.stdout + r.stderr
    assert q(conn, """select count(*) from public.building_permit_statistics b join public.permit_series ps on ps.id = b.series_id
                      join public.data_sources d on d.id = ps.source_id where d.source_key = 'minvu_cedoc_permisos_mensual'""")[0][0] == 346 * 31 * 6
    assert q(conn, "select min(period)::text from public.building_permit_statistics b join public.permit_series ps on ps.id = b.series_id where ps.granularity = 'month'")[0][0] == "2024-01-01"
    ytd = q(conn, "select sum(value), sum(prev_value), bool_and(coverage_ok) from public.get_permit_choropleth('minvu_cedoc_permisos_mensual', 'ytd', 2026, 7)")[0]
    assert ytd[0] is not None and ytd[2] is True
    meses = [r[0] for r in q(conn, """select sum(value) from public.building_permit_statistics b join public.permit_series ps on ps.id = b.series_id
                                       where ps.source_sheet = 'número_total' and ps.granularity = 'month' and b.period >= '2026-01-01' group by b.period order by b.period""")]
    assert meses == [5981, 5961, 6752, 6298, 5917, 4811, 6482]


def test_un_cuadre_roto_no_deja_nada(conn, tmp_path):
    before = q(conn, "select count(*) from public.source_snapshots")[0][0]
    wb = openpyxl.load_workbook(ANUAL)
    wb["número_total"].cell(8, 3).value += 1
    f = tmp_path / "roto.xlsx"
    wb.save(f)
    r = cli("permisos_minvu.py", ["--file", str(f), "--origin-url", "https://example.test/roto.xlsx", "--no-storage"])
    assert r.returncode == 3
    assert q(conn, "select count(*) from public.source_snapshots")[0][0] == before
    assert q(conn, "select count(*) from public.permit_stage")[0][0] == 0
