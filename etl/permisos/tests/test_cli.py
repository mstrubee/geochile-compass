"""CLI etl/permisos_minvu.py: el dry-run no escribe NADA y los casos de error devuelven el código documentado."""
import json
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from conftest import ANUAL, MENSUAL

CLI = Path(__file__).resolve().parents[2] / "permisos_minvu.py"


def run(args, cwd, env_extra=None):
    env = {k: v for k, v in os.environ.items() if k not in ("SUPABASE_DB_URL", "SUPABASE_URL", "SUPABASE_SERVICE_ROLE_KEY")}
    env.update(env_extra or {})
    return subprocess.run([sys.executable, "-B", str(CLI), *args], cwd=cwd, capture_output=True, text=True, env=env)


def listing(root):
    return sorted(str(p.relative_to(root)) for p in Path(root).rglob("*"))


def test_dry_run_no_escribe_nada_y_cuadra(tmp_path):
    before_repo = listing(CLI.parent)
    r = run(["--file", str(ANUAL), "--dry-run", "--compare-with", str(MENSUAL), "--declared-date", "2026-09-11"], tmp_path)
    assert r.returncode == 0, r.stderr + r.stdout
    out = r.stdout
    for needle in ("DRY-RUN (no se escribe NADA", "367  →  346 CUT únicos", "21 de Ñuble fusionadas", "alias explícitos . 4",
                   "definitivo hasta . 2024", "133.914", "150.491", "66.947", "42.202", "CRUCE ANUAL ↔ MENSUAL", "LISTO PARA CARGAR"):
        assert needle in out, needle
    assert listing(tmp_path) == []                              # ni un archivo en el directorio de trabajo
    assert listing(CLI.parent) == before_repo                   # ni __pycache__ ni reportes en el repo


def test_dry_run_mensual_y_json(tmp_path):
    r = run(["--file", str(MENSUAL), "--dry-run", "--json"], tmp_path)
    assert r.returncode == 0, r.stderr
    j = json.loads(r.stdout)
    assert j["kind"] == "monthly" and j["unique_cut"] == 346 and j["nuble_merges"] == 21 and j["ready"] is True
    assert j["plan"]["load_from"] == "2021-08-01" and j["controls"]["ok"] == j["controls"]["total"] == 330
    assert j["values_to_stage"] == 346 * 31 * 6


def _tamper(src, dst, fn):
    wb = openpyxl.load_workbook(src)
    fn(wb)
    wb.save(dst)


def test_total_pais_alterado_devuelve_3(tmp_path):
    dst = tmp_path / "t.xlsx"
    def f(wb):
        ws = wb["número_total"]
        ws.cell(8, 3).value = ws.cell(8, 3).value + 1
    _tamper(ANUAL, dst, f)
    r = run(["--file", str(dst), "--dry-run"], tmp_path)
    assert r.returncode == 3 and "DETENIDO" in r.stdout and "✘" in r.stdout


def test_nombre_sin_match_devuelve_2(tmp_path):
    dst = tmp_path / "n.xlsx"
    def f(wb):
        for ws in wb.worksheets:
            ws.cell(27, 2).value = "Arikaa"
    _tamper(ANUAL, dst, f)
    r = run(["--file", str(dst), "--dry-run"], tmp_path)
    assert r.returncode == 2 and "Arikaa" in r.stdout and "DETENIDO" in r.stdout


def test_celda_rechazada_devuelve_5_o_3(tmp_path):
    dst = tmp_path / "c.xlsx"
    def f(wb):
        wb["número_total"].cell(27, 3).value = "s/d"
    _tamper(ANUAL, dst, f)
    r = run(["--file", str(dst), "--dry-run"], tmp_path)
    assert r.returncode in (3, 5) and "non_numeric" in r.stdout


def test_sin_origin_url_no_carga(tmp_path):
    r = run(["--file", str(ANUAL)], tmp_path)
    assert r.returncode == 2 and "--origin-url" in r.stderr


def test_origin_url_debe_ser_https(tmp_path):
    r = run(["--file", str(ANUAL), "--origin-url", "http://x"], tmp_path)
    assert r.returncode == 2


def test_archivo_inexistente(tmp_path):
    assert run(["--file", str(tmp_path / "no.xlsx"), "--dry-run"], tmp_path).returncode == 2


def test_compare_with_exige_el_otro_tipo(tmp_path):
    r = run(["--file", str(ANUAL), "--dry-run", "--compare-with", str(ANUAL)], tmp_path)
    assert r.returncode == 2
